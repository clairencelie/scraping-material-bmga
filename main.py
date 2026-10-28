"""Scraper harga material bangunan: Mitra10, Depo Bangunan, Tokopedia, Shopee.

Contoh:
    python main.py -k "semen 50kg" -k "bata ringan" --pages 2
    python main.py --keywords-file keywords.txt --sites tokopedia depobangunan -o hasil.xlsx
    python main.py -f keywords.txt -o output/harga.xlsx --deskripsi
    python main.py deskripsi output/harga.xlsx        # isi deskripsi ke file hasil lama
    python main.py shopee-login
    python main.py mitra10-stores
"""
from __future__ import annotations

import argparse
import csv
import logging
import sys
import time
from datetime import datetime
from pathlib import Path

from scraper import description, progress, relevance
from scraper.checkpoint import Checkpoint, SettingsChanged
from scraper.depobangunan import DepoBangunanScraper
from scraper.mitra10 import DEFAULT_STORE_ID, Mitra10Scraper, list_stores
from scraper.models import Product
from scraper.shopee import ShopeeScraper, interactive_login
from scraper.tokopedia import SORT_OPTIONS, TokopediaScraper
from scraper.utils import polite_sleep

log = logging.getLogger("scraper")

SITES = ["mitra10", "depobangunan", "tokopedia", "shopee"]


def build_scraper(site: str, args):
    if site == "mitra10":
        return Mitra10Scraper(headless=not args.show_browser, store_id=args.mitra10_store)
    if site == "depobangunan":
        return DepoBangunanScraper()
    if site == "tokopedia":
        return TokopediaScraper(sort=args.sort)
    if site == "shopee":
        return ShopeeScraper(headless=False)
    raise ValueError(site)


def load_keywords(args) -> list[str]:
    keywords = list(args.keyword or [])
    if args.keywords_file:
        for line in Path(args.keywords_file).read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if line and not line.startswith("#"):
                keywords.append(line)
    return keywords


def load_products(path: Path) -> list[Product]:
    """Baca file hasil (.xlsx/.csv). Kolom yang tidak dikenal (mis. ID) diabaikan."""
    fields = set(Product.columns())
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        rows = list(load_workbook(path, read_only=True).active.iter_rows(values_only=True))
        header, records = rows[0], [dict(zip(rows[0], r)) for r in rows[1:]]
    else:
        with path.open(encoding="utf-8-sig", newline="") as f:
            records = list(csv.DictReader(f))
        header = records[0].keys() if records else []

        def conv(k, v):
            if v in ("", None):
                return None
            if k in ("price", "original_price"):
                return int(float(v))
            if k in ("match_score", "discount_percent", "rating"):
                return float(v)
            if k in ("relevant", "in_stock"):
                return v == "True"
            return v

        records = [{k: conv(k, v) for k, v in r.items()} for r in records]
    missing = {"source", "keyword", "name", "price", "url"} - set(header)
    if missing:
        raise SystemExit(f"{path}: bukan file hasil scraper (kolom hilang: {', '.join(sorted(missing))})")
    return [Product(**{k: v for k, v in r.items() if k in fields}) for r in records]


def save(products: list[Product], output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    columns = Product.columns()
    if output.suffix.lower() == ".xlsx":
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "harga"
        ws.append(columns)
        for p in products:
            ws.append([p.to_dict()[c] for c in columns])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.iter_cols(min_row=1, max_row=1):
            if col[0].value in ("price", "original_price"):
                for cell in ws[col[0].column_letter][1:]:
                    cell.number_format = '"Rp"#,##0'
        wb.save(output)
    else:
        with output.open("w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=columns)
            writer.writeheader()
            writer.writerows(p.to_dict() for p in products)


# Excel hasil sementara ditulis ulang paling sering tiap selang ini (detik).
SNAPSHOT_EVERY = 300


class Snapshot:
    """Menulis file Excel/CSV hasil sementara secara berkala."""

    def __init__(self, output: Path, collect) -> None:
        self.output, self.collect, self.last = output, collect, time.time()

    def maybe(self) -> None:
        if time.time() - self.last >= SNAPSHOT_EVERY:
            self.now()

    def now(self) -> list[Product]:
        products = self.collect()
        try:
            save(products, self.output)
            log.info("[progres] hasil sementara disimpan: %s (%d baris)", self.output, len(products))
        except OSError as e:  # mis. file sedang dibuka & dikunci Excel
            log.warning("[progres] gagal menyimpan %s: %s (dicoba lagi nanti)", self.output, e)
        self.last = time.time()
        return products


def run(args) -> int:
    keywords = load_keywords(args)
    if not keywords:
        print("Tidak ada kata kunci. Pakai -k \"semen\" atau --keywords-file keywords.txt", file=sys.stderr)
        return 2

    output = Path(args.output or f"output/harga_material_{datetime.now():%Y%m%d_%H%M%S}.csv")
    log_file = Path(args.log) if args.log else output.with_suffix(".log")
    progress.setup_logging(args.verbose, log_file)
    if not args.output:
        log.warning("Tanpa -o, nama output memakai jam sekarang sehingga proses TIDAK bisa dilanjutkan "
                    "bila berhenti. Pakai -o output/nama.xlsx untuk jalan yang panjang.")
    meta = {"mode": "scraping", "pages": args.pages, "mitra10_store": args.mitra10_store,
            "sort": args.sort, "min_price": args.min_price}
    try:
        cp = Checkpoint(output, meta, fresh=args.ulang)
    except SettingsChanged as e:
        log.error("%s", e)
        return 2
    progress.say(f"Mulai: {len(keywords)} kata kunci x {len(args.sites)} situs -> {output} (log: {log_file})")

    tasks = [(site, kw) for site in args.sites for kw in keywords]
    cp.plan(args.sites, keywords)
    done_before = sum(1 for t in tasks if cp.is_done(*t))
    if done_before:
        progress.say(f"Melanjutkan: {done_before} / {len(tasks)} kata kunci x situs sudah selesai sebelumnya "
                     f"(dilewati), {len(cp.desc)} deskripsi sudah tersimpan")
    snapshot = Snapshot(output, lambda: cp.products_for(tasks))
    cp.write_status("scraping", force=True)

    # Baris progres di terminal, digambar ulang setiap ada kemajuan.
    view = {"phase": "scraping", "current": ""}

    def show() -> None:
        progress.line.set(cp.summary_line(view["phase"], view["current"]))

    cp.on_change = show
    show()

    try:
        for site in args.sites:
            pending = [kw for kw in keywords if not cp.is_done(site, kw)]
            if not pending:
                continue
            try:
                with build_scraper(site, args) as scraper:
                    for i, kw in enumerate(pending):
                        view["current"] = f"{site}: {kw}"
                        show()
                        try:
                            found = scraper.search(kw, max_pages=args.pages)
                        except Exception as e:  # satu kata kunci gagal tidak menghentikan yang lain
                            log.error("[%s] '%s' gagal: %s", site, kw, e)
                            cp.failures.append(f"{site}: {kw}")
                            show()
                        else:
                            if args.min_price is not None:
                                found = [p for p in found if p.price and p.price >= args.min_price]
                            for p in found:
                                p.match_score, p.relevant = relevance.score(p.keyword, p.name)
                            cp.add_task(site, kw, found)
                        n_done = sum(1 for t in tasks if cp.is_done(*t))
                        log.info("[progres] %d / %d kata kunci x situs (%d%%)", n_done, len(tasks),
                                 n_done * 100 // len(tasks))
                        cp.write_status("scraping")
                        snapshot.maybe()
                        if i < len(pending) - 1:
                            polite_sleep()
            except Exception as e:
                log.error("[%s] tidak bisa dijalankan: %s", site, e)
                cp.failures.append(f"{site}: semua yang belum selesai")

        all_products = cp.products_for(tasks)
        if args.deskripsi:
            progress.line.finish()
            progress.say("Scraping selesai, lanjut mengambil deskripsi produk")
            view["phase"] = "deskripsi"
            cp.write_status("deskripsi", force=True)
            description.enrich(all_products, max_per_keyword=args.deskripsi_maks, cache=cp.desc, tracker=cp)
    except KeyboardInterrupt:
        progress.line.finish()
        cp.on_change = None
        products = snapshot.now()
        cp.write_status("dihentikan (Ctrl+C)", force=True)
        cp.close()
        progress.say(f"Dihentikan. {len(products)} baris hasil sementara ada di {output}. "
                     f"Progres: {cp.status_path}. Lanjutkan dengan menjalankan perintah yang sama.")
        return 130

    progress.line.finish()
    cp.on_change = None
    save(all_products, output)
    complete = all(cp.is_done(*t) for t in tasks)
    cp.write_status("selesai" if complete else "selesai sebagian (ada yang gagal)", force=True)
    cp.close()

    progress.say(f"{len(all_products)} produk disimpan ke {output}")
    for site in args.sites:
        site_products = [p for p in all_products if p.source == site]
        n_relevant = sum(1 for p in site_products if p.relevant)
        n_desc = sum(1 for p in site_products if p.description)
        progress.say(f"  {site:<13} {len(site_products):>5}  (relevan: {n_relevant}"
                     + (f", deskripsi: {n_desc}" if args.deskripsi else "") + ")")
    if cp.failures:
        progress.say("Gagal: " + ", ".join(cp.failures))
        progress.say("Jalankan perintah yang sama untuk mencoba ulang yang gagal.")
    progress.say(f"Progres: {cp.status_path} | jurnal: {cp.journal_path} | log: {log_file}")
    return 0 if all_products else 1


def run_deskripsi(argv: list[str]) -> int:
    """`main.py deskripsi <file>`: isi kolom description ke file hasil yang sudah ada."""
    ap = argparse.ArgumentParser(prog="main.py deskripsi",
                                 description="Isi deskripsi produk ke file hasil scraping yang sudah ada")
    ap.add_argument("file", help="file hasil scraper (.xlsx/.csv)")
    ap.add_argument("-o", "--output", help="file tujuan (default: <file>_deskripsi.<ext>; file asli tidak diubah)")
    ap.add_argument("--maks", type=int, default=10,
                    help="maks produk per kata kunci per situs untuk Tokopedia/Depo Bangunan (default 10, 0 = semua)")
    ap.add_argument("--semua-baris", action="store_true",
                    help="ambil juga untuk produk yang relevant=False (default hanya yang relevan)")
    ap.add_argument("--ulang", action="store_true",
                    help="abaikan progres sebelumnya dan mulai dari awal")
    ap.add_argument("--log", help="file log (default: <file tujuan>.log, ditambahkan di akhir)")
    ap.add_argument("-v", "--verbose", action="store_true", help="tampilkan semua pesan log di terminal")
    args = ap.parse_args(argv)

    src = Path(args.file)
    output = Path(args.output or src.with_name(f"{src.stem}_deskripsi{src.suffix}"))
    log_file = Path(args.log) if args.log else output.with_suffix(".log")
    progress.setup_logging(args.verbose, log_file)
    try:
        cp = Checkpoint(output, {"mode": "deskripsi", "source_file": str(src.resolve())}, fresh=args.ulang)
    except SettingsChanged as e:
        log.error("%s", e)
        return 2
    products = load_products(src)
    for p in products:
        if p.relevant is None:  # file lama sebelum ada kolom relevant
            p.match_score, p.relevant = relevance.score(p.keyword, p.name)
        if not p.description and cp.desc.get(p.url):
            p.description = cp.desc[p.url]
    progress.say(f"Mulai: deskripsi untuk {src} -> {output} (log: {log_file})")
    if cp.desc:
        progress.say(f"Melanjutkan: {len(cp.desc)} deskripsi sudah tersimpan dari jalan sebelumnya")
    cp.write_status("deskripsi", force=True)
    cp.on_change = lambda: progress.line.set(cp.summary_line("deskripsi"))

    def collect() -> list[Product]:
        for p in products:
            if not p.description and cp.desc.get(p.url):
                p.description = cp.desc[p.url]
        return products

    try:
        description.enrich(products, max_per_keyword=args.maks, only_relevant=not args.semua_baris,
                           cache=cp.desc, tracker=cp)
    except KeyboardInterrupt:
        progress.line.finish()
        cp.on_change = None
        Snapshot(output, collect).now()
        cp.write_status("dihentikan (Ctrl+C)", force=True)
        cp.close()
        progress.say(f"Dihentikan. Hasil sementara ada di {output}. Progres: {cp.status_path}. "
                     "Lanjutkan dengan menjalankan perintah yang sama.")
        return 130

    progress.line.finish()
    cp.on_change = None
    save(collect(), output)
    cp.write_status("selesai", force=True)
    cp.close()
    progress.say(f"{len(products)} baris disimpan ke {output}")
    for site in dict.fromkeys(p.source for p in products):
        site_products = [p for p in products if p.source == site]
        progress.say(f"  {site:<13} deskripsi terisi {sum(1 for p in site_products if p.description):>5} / {len(site_products)}")
    progress.say(f"Progres: {cp.status_path} | jurnal: {cp.journal_path} | log: {log_file}")
    return 0


def main() -> int:
    if len(sys.argv) > 1 and sys.argv[1] == "deskripsi":
        return run_deskripsi(sys.argv[2:])
    if len(sys.argv) > 1 and sys.argv[1] == "shopee-login":
        interactive_login()
        return 0
    if len(sys.argv) > 1 and sys.argv[1] == "mitra10-stores":
        for st in sorted(list_stores(), key=lambda st: st["vendor_name"]):
            print(f"{st['vendor_id']:>4}  {st['vendor_name']:<32} {st['street']}")
        return 0

    ap = argparse.ArgumentParser(description="Scraper harga material bangunan")
    ap.add_argument("-k", "--keyword", action="append", help="kata kunci (boleh diulang)")
    ap.add_argument("-f", "--keywords-file", help="file berisi satu kata kunci per baris")
    ap.add_argument("-s", "--sites", nargs="+", choices=SITES, default=SITES[:3],
                    help="situs yang di-scrape (default: mitra10 depobangunan tokopedia)")
    ap.add_argument("-p", "--pages", type=int, default=1, help="jumlah halaman per kata kunci (default 1)")
    ap.add_argument("-o", "--output", help="file output .csv atau .xlsx (default: output/harga_material_<waktu>.csv)")
    ap.add_argument("--sort", choices=list(SORT_OPTIONS), default="relevan", help="urutan hasil Tokopedia")
    ap.add_argument("--mitra10-store", type=int, default=DEFAULT_STORE_ID,
                    help=f"ID toko Mitra10 untuk harga (default {DEFAULT_STORE_ID} = Q-BIG); lihat `main.py mitra10-stores`")
    ap.add_argument("--min-price", type=int, help="buang produk di bawah harga ini (mis. aksesoris murah)")
    ap.add_argument("--deskripsi", action="store_true",
                    help="ambil juga deskripsi produk ke kolom description (menambah waktu, lihat --deskripsi-maks)")
    ap.add_argument("--deskripsi-maks", type=int, default=10,
                    help="maks produk relevan per kata kunci per situs yang diambil deskripsinya "
                         "untuk Tokopedia/Depo Bangunan (default 10, 0 = semua); Mitra10 selalu semua")
    ap.add_argument("--ulang", action="store_true",
                    help="abaikan progres sebelumnya untuk output ini dan mulai dari awal "
                         "(default: jalankan ulang perintah yang sama = melanjutkan)")
    ap.add_argument("--show-browser", action="store_true", help="tampilkan jendela browser Mitra10")
    ap.add_argument("--log", help="file log (default: <output>.log di samping file output, ditambahkan di akhir)")
    ap.add_argument("-v", "--verbose", action="store_true", help="tampilkan semua pesan log di terminal")
    args = ap.parse_args()
    return run(args)


if __name__ == "__main__":
    sys.exit(main())
